"""Builds model/ALHC_model.xlsx: a quarterly three-statement model (Q1'26-Q4'28), the 2027 MBR bridge,
a California statutory-capital (TNE) check, and the valuation with peer multiples, sensitivity grid and
price-implied MBR. Every hardcoded input is blue with a source note; everything else is a formula.
Run scripts/recalc (LibreOffice) afterwards so cached values exist, then read_model.py prints them.
"""
import os, csv, io, contextlib, sys
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter as L
HERE = os.path.dirname(os.path.abspath(__file__)); sys.path.insert(0, HERE)
D = os.path.join(HERE, '..', 'data'); OUT = os.path.join(HERE, '..', 'model'); os.makedirs(OUT, exist_ok=True)
with contextlib.redirect_stdout(io.StringIO()):
    import value_v8 as V8
STARS_COST = 0.5 * V8.cost() + 0.25 * V8.cost(margin=0, margin_later=0) + 0.25 * V8.cost(margin=0, margin_later=0, growth_pts=0.10)
STARS28 = 0.5 * 0.77 * (1 - 0.25) * V8.V.L35

F = 'Arial'
fn = lambda **k: Font(name=F, size=10, **k)
BLUE, GREEN, BOLD = fn(color='0000FF'), fn(color='008000'), fn(bold=True)
YEL = PatternFill('solid', fgColor='FFFF00'); HDR = PatternFill('solid', fgColor='DDEBF7')
NUM, PCT, PCT2, MULT, DOL, DOL2, KM = '#,##0.0;(#,##0.0);-', '0.0%;(0.0%);-', '0.00%;(0.00%);-', '0.0x', '$#,##0;($#,##0);-', '$#,##0.00;($#,##0.00);-', '#,##0.0'

wb = Workbook()
# ---------------- Inputs ----------------
I = wb.active; I.title = 'Inputs'
I['A1'] = 'ALHC model inputs (blue = hardcoded input; yellow = key assumption). $ in millions unless noted.'; I['A1'].font = BOLD
inputs = [
 # name, value, fmt, source, key
 ('price', 8.28, DOL2, 'Nasdaq close 16 Oct 2026 (Yahoo chart API; Google Finance)', False),
 ('price_sep14', 12.94, DOL2, 'Close 14 Sep 2026 (S&P Global via StockAnalysis; Yahoo)', False),
 ('price_oct8', 8.71, DOL2, 'Close 8 Oct 2026', False),
 ('basic_shares', 207.435149, NUM, 'Shares outstanding 27 Jul 2026, 10-Q cover (m)', False),
 ('q3_rev_guide_mid', 1310, NUM, 'Q3 2026 revenue guidance $1,300-1,320m (Q2 release)', False),
 ('q3_gp_guide_mid', 153, NUM, 'Q3 2026 adj. gross profit guidance $148-158m (Q2 call)', False),
 ('q3_ebitda_guide_mid', 25, NUM, 'Q3 2026 adj. EBITDA guidance $20-30m', False),
 ('q3_members_mid', 296.5, NUM, 'Q3 2026 membership guidance 295.5-297.5k', False),
 ('fy26_rev_guide_mid', 5210, NUM, 'FY2026 revenue guidance $5,195-5,225m', False),
 ('fy26_gp_guide_mid', 640, NUM, 'FY2026 adj. gross profit guidance $630-650m (Q2 call)', False),
 ('fy26_ebitda_guide_mid', 154, NUM, 'FY2026 adj. EBITDA guidance $145-163m', False),
 ('fy26_members_mid', 299.5, NUM, 'FY2026 membership guidance 298-301k', False),
 ('my_q3_mbr', 0.887, PCT, "Author: Q3 at the low end of gross-profit guidance (costs ran hot Jul-Aug, Baird 15 Sep)", True),
 ('my_q4_mbr', 0.884, PCT, 'Author: Q4 MBR that puts FY2026 adj. EBITDA at ~$145m, the low end of guidance', True),
 ('g27', 0.18, PCT, 'Author: Dec-27 vs Dec-26 membership growth. Management targets ~20% (Q2 call); consensus revenue implies ~17%', True),
 ('g28', 0.15, PCT, 'Author assumption: 2028 membership growth', True),
 ('jan_share', 0.80, PCT, 'Author: share of the year\'s growth that lands on 1 Jan (annual enrollment period)', False),
 ('rate27', 0.0498, PCT2, 'CMS CY2027 Rate Announcement: +4.98% incl. risk score trend (fact sheet 6 Apr 2026)', False),
 ('rate28', 0.03, PCT, 'Author assumption: 2028 revenue PMPM growth (2028 rate notice not out until Apr 2027)', True),
 ('sga27', 0.090, PCT, 'Author: adj. SG&A % revenue 2027 (2026 guidance midpoint ~9.3%)', True),
 ('sga28', 0.088, PCT, 'Author: adj. SG&A % revenue 2028', True),
 ('mbr28_delta', 0.0, PCT, 'Author: 2028 MBR change vs 2027 before the stars hit', True),
 ('stars28', round(STARS28, 1), NUM, 'value_v8.py: expected 2028 stars margin hit, blend of paths (A 77% x 75% x $133.6m, B none)', False),
 ('stars_cost_ps', round(STARS_COST, 3), DOL2, 'value_v8.py: PV of the downgrade per share, blend (A 50%, B 25%/25%)', False),
 ('da_growth', 0.10, PCT, 'Author: D&A growth a year', False),
 ('sbc_growth', 0.05, PCT, 'Author: stock comp growth a year', False),
 ('dil_q', 0.006, PCT2, 'Author: diluted share growth a quarter (SBC dilution)', False),
 ('tax_gaap', 0.03, PCT, 'Effective tax H1 2026 2.6% (10-Q); NOLs shelter federal tax', False),
 ('tax_norm', 0.24, PCT, 'Author: normalized tax rate to compare with tax-paying peers', False),
 ('ar_pct', 296.893 / 1335.633, PCT, 'Receivables / Q2 2026 revenue (10-Q)', False),
 ('prep_pct', 161.845 / 1335.633, PCT, 'Prepaid & other / Q2 2026 revenue (10-Q)', False),
 ('ap_pct', 32.270 / 1335.633, PCT, 'AP & accrued / Q2 2026 revenue (10-Q)', False),
 ('dcp_days', 48.3, NUM, 'Days claims payable Q2 2026 (MEP / quarterly medical x 91; ops_model.py)', False),
 ('capex_q', 9.0, NUM, 'Capex a quarter (H1 2026 $17.5m, 10-Q cash flow)', False),
 ('accretion_q', 0.57, NUM, 'Debt-cost amortization a quarter (H1 2026 $1.15m)', False),
 ('parent_cash', 135.4, NUM, 'Operating parent cash & ST investments 30 Jun 2026 (10-Q); held flat', False),
 ('revolver', 200.0, NUM, 'Citibank senior secured revolver, matures 26 Feb 2029, undrawn 30 Jun 2026 (10-Q)', False),
 ('convert_face', 330.0, NUM, '4.25% convertible notes due 15 Nov 2029', False),
 ('nol_fed', 617.001, NUM, 'Federal NOLs 31 Dec 2025 (10-K)', False),
 ('nol_pv_factor', 0.79, PCT, 'Author: NOLs used over ~4 years, 10% discount', False),
 ('ca_share', 0.78, PCT, 'Author: California share of premium (H3815 was 75% of members, Oct 2026 CMS file; plus H6697/H8832)', True),
 ('cap_share_low', 0.0, PCT, 'TNE check: share of CA medical spend paid on a capitated basis (low case)', False),
 ('cap_share_high', 0.5, PCT, 'TNE check: capitated share (high case); not disclosed', False),
 # bridge
 ('egr27', 0.0533, PCT2, 'CMS CY2027 effective growth rate (Original Medicare per-capita cost trend), fact sheet', False),
 ('retention', 0.88, PCT, 'Author: annual member retention (H3815 voluntary disenrollment 10% plus deaths)', False),
 ('g26_growth', 0.315, PCT, 'Membership growth Q2 2026 y/y 31.5% (Q2 release)', False),
 ('gp_now', 640, NUM, 'Adj. gross profit at 2026 guidance midpoint (Q2 call)', False),
 ('gp_embedded', 880, NUM, "Management's embedded gross-profit potential of current book (Q2 call)", False),
 ('maturity_haircut', 0.5, PCT, "Author: credit given to management's maturity estimate in the base case", True),
 ('benefit_change', -0.002, PCT, 'Author: 2027 benefit trims (CMS landscape: MOOP up on 2 of 8 matched LA plans, drug deductible added on 2)', True),
 ('cost_persist', 0.5, PCT, 'Author: share of H2-2026 MBR pressure that persists into 2027 (base)', True),
]
row = {}
I['A3'], I['B3'], I['C3'] = 'Input', 'Value', 'Source / note'
for c in 'ABC': I[f'{c}3'].font = BOLD; I[f'{c}3'].fill = HDR
for k, (name, val, fmt, src, key) in enumerate(inputs):
    r = 4 + k; row[name] = r
    I[f'A{r}'] = name; I[f'A{r}'].font = fn()
    I[f'B{r}'] = val; I[f'B{r}'].font = BLUE; I[f'B{r}'].number_format = fmt
    if key: I[f'B{r}'].fill = YEL
    I[f'C{r}'] = src; I[f'C{r}'].font = fn()
I.column_dimensions['A'].width = 22; I.column_dimensions['B'].width = 12; I.column_dimensions['C'].width = 110
inp = lambda n: f"Inputs!$B${row[n]}"

# ---------------- Bridge ----------------
B = wb.create_sheet('MBR Bridge')
B['A1'] = '2027 medical benefit ratio bridge (MBR = adjusted medical cost / revenue)'; B['A1'].font = BOLD
hdr = ['Step', 'Bear', 'Base', 'Bull', 'How it is built']
for j, h in enumerate(hdr): c = B.cell(3, 1 + j, h); c.font = BOLD; c.fill = HDR
steps = [
 ('2026E MBR (model)', "=Model!$P$10", "=Model!$P$10", "=Model!$P$10", 'My FY2026 estimate (Model sheet); guidance midpoint implies ~87.7%'),
 ('Members in year 1-2, 2026', f"=(1-{inp('retention')}/(1+{inp('g26_growth')}))*(1+{inp('retention')}/(1+{inp('g26_growth')}))", None, None, 'Steady-state share at 31.5% growth, 88% retention: new + second-year members'),
 ('Members in year 1-2, 2027', f"=(1-{inp('retention')}/(1+{inp('g27')}))*(1+{inp('retention')}/(1+{inp('g27')}))", None, None, 'Same at 2027 growth'),
 ('New vs mature MBR gap', f"=(2*({inp('gp_embedded')}-{inp('gp_now')}))/{inp('fy26_rev_guide_mid')}", None, None, "From management's $880m vs $640m: half the book is new, so new-member GP is lower by 2x(880-640) per the whole book"),
 ('Maturity (mix) effect', "=0", f"=-(B5-B6)*B7*{inp('maturity_haircut')}", "=-(B5-B6)*B7", 'Fewer early-tenure members as growth slows; bear gives no credit, base half, bull full'),
 ('Rate vs cost trend', f"=B4*((1+{inp('egr27')}+0.006)/(1+{inp('rate27')})-1)", f"=B4*((1+{inp('egr27')})/(1+{inp('rate27')})-1)", f"=B4*((1+{inp('egr27')})/(1+{inp('rate27')})-1)", 'Cost per member grows with CMS effective growth rate 5.33% vs revenue +4.98%; bear adds 0.6 pts of trend'),
 ('2027 benefit design', "=0", f"={inp('benefit_change')}", f"={inp('benefit_change')}", 'Bids filed June 2026: modest trims in the 2027 plan data'),
 ('H2-2026 cost pressure carried', "=(Model!$S$10-Model!$P$10)", f"=(Model!$S$10-Model!$P$10)*{inp('cost_persist')}", "=0", 'H2-2026 MBR above full year: all of it persists (bear), half (base), none (bull)'),
]
for i, (lab, b, c, d, note) in enumerate(steps):
    r = 4 + i
    B.cell(r, 1, lab).font = fn()
    if c is None:
        B.cell(r, 2, b); B.cell(r, 3, '=B%d' % r); B.cell(r, 4, '=B%d' % r)
    else:
        B.cell(r, 2, b); B.cell(r, 3, c); B.cell(r, 4, d)
    for j in (2, 3, 4): B.cell(r, j).number_format = PCT2; B.cell(r, j).font = fn()
    B.cell(r, 5, note).font = fn()
rT = 4 + len(steps)
B.cell(rT, 1, '2027E MBR').font = BOLD
for j, col in ((2, 'B'), (3, 'C'), (4, 'D')):
    B.cell(rT, j, f"={col}4+{col}8+{col}9+{col}10+{col}11"); B.cell(rT, j).number_format = PCT2; B.cell(rT, j).font = BOLD
B.column_dimensions['A'].width = 30; B.column_dimensions['E'].width = 100
for c in 'BCD': B.column_dimensions[c].width = 11
MBR27_BASE = f"'MBR Bridge'!$C${rT}"

# ---------------- Model ----------------
M = wb.create_sheet('Model')
qs = ["Q1'26A", "Q2'26A", "Q3'26E", "Q4'26E", "Q1'27E", "Q2'27E", "Q3'27E", "Q4'27E", "Q1'28E", "Q2'28E", "Q3'28E", "Q4'28E"]
cols = [L(3 + i) for i in range(12)]  # C..N
ann = {'FY26': 'P', 'FY27': 'Q', 'FY28': 'R', 'H2-26': 'S'}
M['A1'] = 'Alignment Healthcare quarterly model ($m unless noted; members in thousands)'; M['A1'].font = BOLD
M['A3'] = 'Line'; M['B3'] = 'Note'
for i, q in enumerate(qs): M[f'{cols[i]}3'] = q
for k, c in ann.items(): M[f'{c}3'] = k
for c in ['A', 'B'] + cols + list(ann.values()): M[f'{c}3'].font = BOLD; M[f'{c}3'].fill = HDR
def setrow(r, label, vals, fmt=NUM, note='', annual=None, bold=False):
    M[f'A{r}'] = label; M[f'A{r}'].font = BOLD if bold else fn(); M[f'B{r}'] = note; M[f'B{r}'].font = fn()
    for i, v in enumerate(vals):
        if v is None: continue
        c = M[f'{cols[i]}{r}']; c.value = v; c.number_format = fmt
        c.font = BLUE if not (isinstance(v, str) and v.startswith('=')) else (BOLD if bold else fn())
    if annual:
        for k, f in annual.items():
            c = M[f'{ann[k]}{r}']; c.value = f; c.number_format = fmt; c.font = BOLD if bold else fn()
SUMA = lambda r: {'FY26': f'=SUM(C{r}:F{r})', 'FY27': f'=SUM(G{r}:J{r})', 'FY28': f'=SUM(K{r}:N{r})', 'H2-26': f'=SUM(E{r}:F{r})'}
ENDA = lambda r: {'FY26': f'=F{r}', 'FY27': f'=J{r}', 'FY28': f'=N{r}'}
# membership
r = 5
setrow(r, 'Ending members (k)', [None, 294.1, f"={inp('q3_members_mid')}", f"={inp('fy26_members_mid')}",
       f"=F5*(1+{inp('g27')}*{inp('jan_share')})", f"=G5+(F5*(1+{inp('g27')})-G5)/3", f"=H5+(F5*(1+{inp('g27')})-G5)/3", f"=F5*(1+{inp('g27')})",
       f"=J5*(1+{inp('g28')}*{inp('jan_share')})", f"=K5+(J5*(1+{inp('g28')})-K5)/3", f"=L5+(J5*(1+{inp('g28')})-K5)/3", f"=J5*(1+{inp('g28')})"],
       KM, 'Q2 actual 294,100 (release); Q3/Q4 guidance midpoints; 2027-28 growth inputs', ENDA(5))
setrow(6, 'Average members (k)', [None, None] + [(f'={cols[i]}5' if i in (4, 8) else f'=({cols[i-1]}5+{cols[i]}5)/2') for i in range(2, 12)], KM, 'Average of quarter ends; Q1 = quarter end because the January step lands on 1 Jan')
setrow(7, 'Revenue PMPM ($)', [None, None, '=E8/(E6*3)*1000', '=F8/(F6*3)*1000'] + [f"=AVERAGE($E$7:$F$7)*(1+{inp('rate27')})"] * 4 + [f"=$G$7*(1+{inp('rate28')})"] * 4,
       '#,##0', '2026 from guidance; 2027 = H2-26 PMPM x (1 + 4.98%); 2028 = 2027 x (1 + rate28)')
setrow(8, 'Revenue', [1235.197, 1335.633, f"={inp('q3_rev_guide_mid')}", f"={inp('fy26_rev_guide_mid')}-C8-D8-E8"] + [f'={c}6*3*{c}7/1000' for c in cols[4:]],
       NUM, 'Q1/Q2 actual (10-Q); Q3 guidance midpoint; Q4 = FY guide less 9M', SUMA(8), bold=True)
setrow(9, 'Adj. gross profit (actual/guide)', [146.0, 183.0, None, None] + [None] * 8, NUM, 'Q1 = H1 approx $329m less Q2 $183m (Q2 call); Q2 reported')
setrow(10, 'Adj. MBR', ['=1-C9/C8', '=1-D9/D8', f"={inp('my_q3_mbr')}", f"={inp('my_q4_mbr')}"] +
       [f"={cols[i-4]}10-$P$10+{MBR27_BASE}" for i in range(4, 8)] + [f"={cols[i-4]}10-$Q$10+$Q$10+{inp('mbr28_delta')}" for i in range(8, 12)],
       PCT, '2027 quarters keep the 2026 seasonal shape around the bridge result; 2028 = 2027 + delta',
       {'FY26': '=1-P12/P8', 'FY27': '=1-Q12/Q8', 'FY28': '=1-R12/R8', 'H2-26': '=1-S12/S8'})
setrow(11, 'Adj. medical cost', [f'={c}8-{c}9' for c in cols[:2]] + [f'={c}8*{c}10' for c in cols[2:]], NUM, '', SUMA(11))
setrow(12, 'Adj. gross profit', [f'={c}8-{c}11' for c in cols], NUM, '', SUMA(12), bold=True)
setrow(13, 'Adj. SG&A', ['=C12-C14', '=D12-D14', f"={inp('q3_gp_guide_mid')}-{inp('q3_ebitda_guide_mid')}",
       f"=({inp('fy26_gp_guide_mid')}-{inp('fy26_ebitda_guide_mid')})-C13-D13-E13"] + [f"={c}8*{inp('sga27')}" for c in cols[4:8]] + [f"={c}8*{inp('sga28')}" for c in cols[8:]],
       NUM, 'Actual = GP less EBITDA; Q3 guide-implied; Q4 = FY guide-implied less 9M; then % of revenue', SUMA(13))
setrow(14, 'Adj. EBITDA before stars', [37.851, 68.145] + [f'={c}12-{c}13' for c in cols[2:]], NUM, 'Q1/Q2 reported (Q2 release: H1 $106.0m, Q2 $68.1m)', SUMA(14), bold=True)
setrow(15, 'Expected 2028 stars hit', [0] * 8 + [f"=-{inp('stars28')}/4"] * 4, NUM, 'Blend of paths, lawsuit odds 25% (value_v8.py)', SUMA(15))
setrow(16, 'Adj. EBITDA', [f'={c}14+{c}15' for c in cols], NUM, '', SUMA(16), bold=True)
setrow(17, 'D&A', [7.839, 7.854, '=D17', '=D17'] + [f"=$F$17*(1+{inp('da_growth')})"] * 4 + [f"=$G$17*(1+{inp('da_growth')})"] * 4, NUM, 'Q1/Q2 10-Q', SUMA(17))
setrow(18, 'Interest expense', [4.062, 4.275] + ['=$D$18'] * 10, NUM, 'Q1/Q2 10-Q; revolver undrawn', SUMA(18))
setrow(19, 'Stock-based compensation', [14.019, 18.174, '=D19', '=D19'] + [f"=$D$19*(1+{inp('sbc_growth')})"] * 4 + [f"=$G$19*(1+{inp('sbc_growth')})"] * 4, NUM, 'Q1/Q2 10-Q (H1 $32.2m)', SUMA(19))
setrow(20, 'Pre-tax income (approx. GAAP)', [11.441, 37.826] + [f'={c}16-{c}17-{c}18-{c}19' for c in cols[2:]], NUM, 'Q1/Q2 reported pre-tax (10-Q)', SUMA(20))
setrow(21, 'Income tax', [0.025, 1.266] + [f"=MAX(0,{c}20)*{inp('tax_gaap')}" for c in cols[2:]], NUM, '', SUMA(21))
setrow(22, 'Net income', [f'={c}20-{c}21' for c in cols], NUM, '', SUMA(22), bold=True)
setrow(23, 'Adj. net income (EBITDA - D&A - interest)', [f'={c}16-{c}17-{c}18' for c in cols], NUM, 'Matches the consensus adj. EPS definition used in ops_model.py', SUMA(23))
setrow(24, 'Diluted shares (m)', [213.7, 215.024] + [f"={cols[i-1]}24*(1+{inp('dil_q')})" for i in range(2, 12)], NUM, 'Q2 10-Q 215.0m; Q1 approx.', {'FY26': '=AVERAGE(C24:F24)', 'FY27': '=AVERAGE(G24:J24)', 'FY28': '=AVERAGE(K24:N24)'})
setrow(25, 'Adj. EPS ($)', [f'={c}23/{c}24' for c in cols], DOL2, '', {'FY26': '=P23/P24', 'FY27': '=Q23/Q24', 'FY28': '=R23/R24'}, bold=True)
setrow(26, 'GAAP EPS ($)', [f'={c}22/{c}24' for c in cols], DOL2, '', {'FY26': '=P22/P24', 'FY27': '=Q22/Q24', 'FY28': '=R22/R24'})
# balance sheet
M['A28'] = 'Balance sheet (quarter end)'; M['A28'].font = BOLD
fc = cols[2:]  # forecast columns E..N
prev = lambda c: L(ord(c) - 64 - 1)
bs = [
 (29, 'Cash and equivalents', 693.453, lambda c: f'={prev(c)}29+{c}63', '30 Jun 2026 10-Q; then cash flow'),
 (30, 'Accounts receivable', 296.893, lambda c: f"={c}8*{inp('ar_pct')}", ''),
 (31, 'Investments, current', 8.258, lambda c: f'={prev(c)}31', ''),
 (32, 'Prepaid and other current', 161.845, lambda c: f"={c}8*{inp('prep_pct')}", ''),
 (33, 'Property and equipment', 66.211, lambda c: f"={prev(c)}33+{inp('capex_q')}-{c}17", ''),
 (34, 'Right-of-use, goodwill, intangibles, other', 7.304 + 32.060 + 4.550 + 8.464, lambda c: f'={prev(c)}34', ''),
 (35, 'Total assets', None, lambda c: f'=SUM({c}29:{c}34)', ''),
 (37, 'Medical expenses payable', 612.748, lambda c: f"={c}11/91*{inp('dcp_days')}", 'Days claims payable held at 48.3'),
 (38, 'Accounts payable and accrued', 32.270, lambda c: f"={c}8*{inp('ap_pct')}", ''),
 (39, 'Accrued compensation and leases', 38.389 + 6.444, lambda c: f'={prev(c)}39', ''),
 (40, 'Convertible notes (carrying)', 324.056, lambda c: f"=MIN({inp('convert_face')},{prev(c)}40+{inp('accretion_q')})", 'Accretes to $330m face'),
 (41, 'Total liabilities', None, lambda c: f'=SUM({c}37:{c}40)', ''),
 (42, 'Stockholders\' equity', 265.131, lambda c: f'={prev(c)}42+{c}22+{c}19', 'Prior + net income + stock comp'),
 (43, 'Total liabilities and equity', None, lambda c: f'={c}41+{c}42', ''),
 (44, 'Check (assets - L&E)', None, lambda c: f'=ROUND({c}35-{c}43,6)', 'Must be zero'),
]
for r, lab, d0, f, note in bs:
    M[f'A{r}'] = lab; M[f'A{r}'].font = BOLD if 'Total' in lab or 'Check' in lab else fn(); M[f'B{r}'] = note; M[f'B{r}'].font = fn()
    if d0 is not None:
        M[f'D{r}'] = d0; M[f'D{r}'].font = BLUE; M[f'D{r}'].number_format = NUM
    else:
        M[f'D{r}'] = f(cols[1]).replace(cols[1] + '63', '0'); M[f'D{r}'].number_format = NUM
    for c in fc:
        M[f'{c}{r}'] = f(c); M[f'{c}{r}'].number_format = NUM; M[f'{c}{r}'].font = fn()
    for k, c in (('FY26', 'F'), ('FY27', 'J'), ('FY28', 'N')):
        M[f'{ann[k]}{r}'] = f'={c}{r}'; M[f'{ann[k]}{r}'].number_format = NUM
# cash flow
M['A46'] = 'Cash flow (quarter)'; M['A46'].font = BOLD
cf = [
 (47, 'Net income', lambda c: f'={c}22'),
 (48, 'D&A', lambda c: f'={c}17'),
 (49, 'Stock-based compensation', lambda c: f'={c}19'),
 (50, 'Debt-cost amortization', lambda c: f'={c}40-{prev(c)}40'),
 (51, 'Change in receivables', lambda c: f'=-({c}30-{prev(c)}30)'),
 (52, 'Change in prepaid and other', lambda c: f'=-({c}32-{prev(c)}32)'),
 (53, 'Change in medical expenses payable', lambda c: f'={c}37-{prev(c)}37'),
 (54, 'Change in payables and accrued', lambda c: f'={c}38-{prev(c)}38'),
 (55, 'Cash from operations', lambda c: f'=SUM({c}47:{c}54)'),
 (57, 'Capital expenditure', lambda c: f"=-{inp('capex_q')}"),
 (58, 'Cash from investing', lambda c: f'={c}57'),
 (60, 'Cash from financing', lambda c: '=0'),
 (63, 'Net change in cash', lambda c: f'={c}55+{c}58+{c}60'),
]
for r, lab, f in cf:
    M[f'A{r}'] = lab; M[f'A{r}'].font = BOLD if 'Cash' in lab or 'Net change' in lab else fn()
    for c in fc:
        M[f'{c}{r}'] = f(c); M[f'{c}{r}'].number_format = NUM; M[f'{c}{r}'].font = fn()
    M[f'S{r}'] = f'=SUM(E{r}:F{r})'; M[f'Q{r}'] = f'=SUM(G{r}:J{r})'; M[f'R{r}'] = f'=SUM(K{r}:N{r})'
    for c in 'QRS': M[f'{c}{r}'].number_format = NUM
M['B55'] = 'H1 2026 actual CFO $111.4m (10-Q); FY26 column shows H2 only in col S'; M['B55'].font = fn()
M.column_dimensions['A'].width = 40; M.column_dimensions['B'].width = 46
for c in cols + list(ann.values()): M.column_dimensions[c].width = 10
M.freeze_panes = 'C4'

# ---------------- Statutory capital ----------------
S = wb.create_sheet('Statutory Capital')
S['A1'] = 'California tangible net equity (TNE) check, Knox-Keene formula as described in the Q2 2026 10-Q, note 11'; S['A1'].font = BOLD
S['A2'] = 'Required = 150% of the largest of $1m; 2% of first $150m of annualized premium + 1% above; 8% of first $150m of non-capitated health care spend + 4% above. Amounts are not disclosed; this is an outside estimate.'; S['A2'].font = fn()
for j, h in enumerate(['Line', 'FY26E', 'FY27E', 'FY28E', 'Note']):
    c = S.cell(4, 1 + j, h); c.font = BOLD; c.fill = HDR
mc = {'B': 'P', 'C': 'Q', 'D': 'R'}
lines = [
 ('CA annualized premium', lambda m: f"=Model!{m}8*{inp('ca_share')}", 'Revenue x CA share'),
 ('CA medical spend', lambda m: f"=Model!{m}11*{inp('ca_share')}", ''),
 ('Required: premium test', lambda m: None, '2% of first $150m + 1% above'),
 ('Required: spend test, 0% capitated', lambda m: None, '8% / 4%, no capitation'),
 ('Required: spend test, 50% capitated', lambda m: None, '8% / 4% on half the spend'),
 ('Required TNE (150%), low case', lambda m: None, 'Max of premium test and 50%-capitated spend test'),
 ('Required TNE (150%), high case', lambda m: None, 'Max of premium test and 0%-capitated spend test'),
 ('Estimated regulated net worth', lambda m: f"=Model!{m}42+Model!{m}40-{inp('parent_cash')}-(32.06+4.55)", 'Equity + parent debt - parent cash - goodwill/intangibles (rough)'),
 ('Headroom, low requirement', lambda m: None, ''),
 ('Headroom, high requirement', lambda m: None, ''),
]
for i, (lab, f, note) in enumerate(lines):
    r = 5 + i; S.cell(r, 1, lab).font = fn(); S.cell(r, 5, note).font = fn()
    for col, m in mc.items():
        if i == 2: v = f"=0.02*MIN(150,{col}5)+0.01*MAX(0,{col}5-150)"
        elif i == 3: v = f"=0.08*MIN(150,{col}6*(1-{inp('cap_share_low')}))+0.04*MAX(0,{col}6*(1-{inp('cap_share_low')})-150)"
        elif i == 4: v = f"=0.08*MIN(150,{col}6*(1-{inp('cap_share_high')}))+0.04*MAX(0,{col}6*(1-{inp('cap_share_high')})-150)"
        elif i == 5: v = f"=1.5*MAX(1,{col}7,{col}9)"
        elif i == 6: v = f"=1.5*MAX(1,{col}7,{col}8)"
        elif i == 8: v = f"={col}12-{col}10"
        elif i == 9: v = f"={col}12-{col}11"
        else: v = f(m)
        S[f'{col}{r}'] = v; S[f'{col}{r}'].number_format = NUM; S[f'{col}{r}'].font = fn()
S.column_dimensions['A'].width = 40; S.column_dimensions['E'].width = 60

# ---------------- Valuation ----------------
Vs = wb.create_sheet('Valuation')
Vs['A1'] = 'Valuation: same method for any managed-care stock. Peer data Yahoo Finance, 16 Oct 2026 closes and 17 Oct consensus.'; Vs['A1'].font = BOLD
ph = ['Ticker', 'Price 16 Oct', 'Price 14 Sep', 'Shares (m)', 'EV ($m)', 'EBITDA TTM ($m)', 'EPS 2026E', 'EPS 2027E', 'Fwd P/E', 'Fwd EBITDA 2027 ($m)', 'Fwd EV/EBITDA', 'P/E at 14 Sep', 'EV/EBITDA at 14 Sep']
for j, h in enumerate(ph): c = Vs.cell(3, 1 + j, h); c.font = BOLD; c.fill = HDR
peers = list(csv.DictReader(open(os.path.join(D, 'peers_2026_10_16.csv'))))
for i, p in enumerate(peers):
    r = 4 + i
    vals = [p['ticker'], float(p['price_2026_10_16']), float(p['close_2026_09_14']), float(p['shares_out']) / 1e6, float(p['ev_usd']) / 1e6, float(p['ebitda_ttm_usd']) / 1e6, float(p['eps_2026_cons']), float(p['eps_2027_cons'])]
    for j, v in enumerate(vals):
        c = Vs.cell(r, 1 + j, v); c.font = BLUE if j else fn(); c.number_format = '#,##0.00' if j else '@'
    Vs.cell(r, 9, f'=B{r}/H{r}').number_format = MULT
    Vs.cell(r, 10, f"=F{r}+(H{r}-G{r})*D{r}/(1-{inp('tax_norm')})").number_format = '#,##0'
    Vs.cell(r, 11, f'=E{r}/J{r}').number_format = MULT
    Vs.cell(r, 12, f'=C{r}/H{r}').number_format = MULT
    Vs.cell(r, 13, f'=(E{r}+(C{r}-B{r})*D{r})/J{r}').number_format = MULT
n = len(peers); rM = 4 + n
Vs.cell(rM, 1, 'Median, all seven').font = BOLD
Vs.cell(rM + 1, 1, 'Median ex HUM, MOH (trough 2026 EPS)').font = BOLD
for col in 'IKLM':
    Vs[f'{col}{rM}'] = f'=MEDIAN({col}4:{col}{rM-1})'; Vs[f'{col}{rM}'].number_format = MULT; Vs[f'{col}{rM}'].font = BOLD
    keep = [f'{col}{4+i}' for i, p in enumerate(peers) if p['ticker'] not in ('HUM', 'MOH')]
    Vs[f'{col}{rM+1}'] = f'=MEDIAN({",".join(keep)})'; Vs[f'{col}{rM+1}'].number_format = MULT
Vs.cell(rM + 2, 1, 'Fwd EBITDA 2027 = TTM EBITDA + (EPS27 - EPS26) x shares / (1 - 24%): an approximation; Yahoo does not publish forward EBITDA. Sep 14 multiples use today\'s consensus EPS.').font = fn()
# ALHC target
r0 = rM + 4
Vs.cell(r0, 1, 'ALHC price target (house method)').font = BOLD
tl = [
 ('2027E revenue', '=Model!Q8', NUM), ('2027E MBR', '=Model!Q10', PCT2), ('2027E adj. SG&A ratio', f"={inp('sga27')}", PCT),
 ('2027E adj. EBITDA', '=Model!Q16', NUM), ('2027E D&A + interest', '=Model!Q17+Model!Q18', NUM), ('2027E adj. net income', '=Model!Q23', NUM),
 ('Diluted shares 2027 (m)', '=Model!Q24', NUM), ('Peer forward P/E (median)', f'=I{rM}', MULT), ('Peer forward EV/EBITDA (median)', f'=K{rM}', MULT),
 ('NOL value per share', f"={inp('nol_fed')}*0.21*{inp('nol_pv_factor')}/B{r0+7}", DOL2), ('Parent net debt', f"={inp('convert_face')}-{inp('parent_cash')}", NUM),
 ('P/E value', f"=B{r0+8}*B{r0+6}*(1-{inp('tax_norm')})/B{r0+7}+B{r0+10}", DOL2),
 ('EV/EBITDA value', f'=(B{r0+9}*B{r0+4}-B{r0+11})/B{r0+7}', DOL2),
 ('Less stars cost', f"=-{inp('stars_cost_ps')}", DOL2),
 ('12-month price target', f'=AVERAGE(B{r0+12},B{r0+13})+B{r0+14}', DOL2),
 ('Upside / (downside)', f"=B{r0+15}/{inp('price')}-1", PCT),
]
for i, (lab, f, fmt) in enumerate(tl):
    r = r0 + 1 + i; Vs.cell(r, 1, lab).font = BOLD if 'target' in lab else fn(); c = Vs.cell(r, 2, f); c.number_format = fmt; c.font = BOLD if 'target' in lab else fn()
# implied MBR: target is linear in EBITDA -> solve
ri = r0 + 18
Vs.cell(ri, 1, 'Price-implied 2027 MBR').font = BOLD
def implied(price_ref, pe_ref, eve_ref, stars_ref, row):
    # target = 0.5*[PE*(E-DI)*(1-t)/sh + NOL] + 0.5*[(EVE*E - ND)/sh] - stars  -> solve E
    sh, DI, NOL, ND, REV, SGA = f'B{r0+7}', f'B{r0+5}', f'B{r0+10}', f'B{r0+11}', f'B{r0+1}', f'B{r0+3}'
    t = inp('tax_norm')
    E = f"((2*({price_ref}+{stars_ref})-NOLX+{pe_ref}*{DI}*(1-{t})/{sh}+{ND}/{sh})/({pe_ref}*(1-{t})/{sh}+{eve_ref}/{sh}))".replace('NOLX', NOL)
    # stars28 is inside 2027? No: 2027 EBITDA excludes the 2028 hit, so E maps directly to MBR
    return f"=1-{SGA}-{E}/{REV}"
rows_imp = [
 ('Now ($8.28, full stars cost, today\'s multiples)', implied(inp('price'), f'I{rM}', f'K{rM}', inp('stars_cost_ps'), 0)),
 ('14 Sep ($12.94, 31% of stars cost, 14 Sep multiples)', implied(inp('price_sep14'), f'L{rM}', f'M{rM}', f"0.31*{inp('stars_cost_ps')}", 1)),
]
for i, (lab, f) in enumerate(rows_imp):
    Vs.cell(ri + 1 + i, 1, lab).font = fn(); c = Vs.cell(ri + 1 + i, 2, f); c.number_format = PCT2
Vs.cell(ri + 3, 1, 'Change (points of MBR)').font = BOLD; Vs.cell(ri + 3, 2, f'=B{ri+1}-B{ri+2}').number_format = PCT2
# sensitivity grid: P/E x MBR
rg = ri + 6
Vs.cell(rg, 1, 'Target sensitivity: forward P/E (rows) x 2027 MBR (columns); EV/EBITDA leg at the peer median').font = BOLD
mbrs = [0.870, 0.875, 0.879, 0.885, 0.890, 0.895]; pes = [12.0, 13.0, 14.5, 16.0, 17.0]
for j, m in enumerate(mbrs):
    c = Vs.cell(rg + 1, 2 + j, m); c.number_format = PCT; c.font = BLUE
for i, pe in enumerate(pes):
    c = Vs.cell(rg + 2 + i, 1, pe); c.number_format = MULT; c.font = BLUE
    for j in range(len(mbrs)):
        mc_ = f'{L(2+j)}${rg+1}'; pr = f'$A{rg+2+i}'
        E = f"($B${r0+1}*(1-{mc_}-$B${r0+3}))"
        f = f"=AVERAGE({pr}*({E}-$B${r0+5})*(1-{inp('tax_norm')})/$B${r0+7}+$B${r0+10},($K${rM}*{E}-$B${r0+11})/$B${r0+7})-{inp('stars_cost_ps')}"
        c = Vs.cell(rg + 2 + i, 2 + j, f); c.number_format = DOL2
# scenarios from bridge
rs = rg + 9
Vs.cell(rs, 1, 'Scenarios (MBR from the bridge; all else as the model)').font = BOLD
for j, h in enumerate(['Scenario', '2027 MBR', '2027 adj. EBITDA', '2027 adj. EPS', 'Target']):
    c = Vs.cell(rs + 1, 1 + j, h); c.font = BOLD; c.fill = HDR
for i, (nm, bc) in enumerate((('Bear', 'B'), ('Base', 'C'), ('Bull', 'D'))):
    r = rs + 2 + i; Vs.cell(r, 1, nm)
    Vs.cell(r, 2, f"='MBR Bridge'!{bc}{rT}").number_format = PCT2
    E = f"($B${r0+1}*(1-B{r}-$B${r0+3}))"
    Vs.cell(r, 3, f"={E}").number_format = NUM
    Vs.cell(r, 4, f"=({E}-$B${r0+5})/$B${r0+7}").number_format = DOL2
    Vs.cell(r, 5, f"=AVERAGE($B${r0+8}*({E}-$B${r0+5})*(1-{inp('tax_norm')})/$B${r0+7}+$B${r0+10},($B${r0+9}*{E}-$B${r0+11})/$B${r0+7})-{inp('stars_cost_ps')}").number_format = DOL2
Vs.column_dimensions['A'].width = 48
for j in range(2, 14): Vs.column_dimensions[L(j)].width = 13
for ws in wb.worksheets:
    for rw in ws.iter_rows():
        for c in rw:
            if c.font and c.font.name != F: c.font = Font(name=F, size=10, bold=c.font.bold, color=c.font.color)
path = os.path.join(OUT, 'ALHC_model.xlsx'); wb.save(path); print(path, 'rows', rT, r0, rM, ri, rg, rs)
