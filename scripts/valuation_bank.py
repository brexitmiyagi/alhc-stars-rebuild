"""Price target the way a bank desk would build it, and every way it can be attacked.

Three methods on the same 2027 numbers from model/ALHC_model.xlsx:
  1. Peer forward P/E on 2027 EPS taxed at 24%, plus the NOLs.
  2. Peer forward EV/EBITDA, with consensus 2027 EBITDA for the peers (MarketScreener / S&P Global)
     and the same cash treatment on both sides.
  3. A DCF to 2036 with a terminal margin, and the margin the price implies.

Two things decide the answer and both are shown, not buried:
  - Stock pay. Alignment's adjusted EBITDA and the consensus EPS leave out stock-based compensation
    ($76m in my 2027). Peers' adjusted EPS and the S&P EBITDA for peers include it. Like for like means
    expensing it. I show both.
  - Cash. Peers' net debt nets all their cash, regulated or not. So the consistent treatment for
    Alignment is consolidated cash too. Parent-only and parent plus dividend capacity are shown as
    sensitivities.

The stars cost ($0.937 a share, value_v8.py) comes off each method. The DCF already carries the 2028
part of it inside its cash flows, so only the rest comes off there.
"""
import os, csv
import numpy as np
from openpyxl import load_workbook
from scipy.optimize import brentq

HERE = os.path.dirname(os.path.abspath(__file__))
D = lambda f: os.path.join(HERE, '..', 'data', f)
out = []
P = lambda s='': (print(s), out.append(s))

wb = load_workbook(os.path.join(HERE, '..', 'model', 'ALHC_model.xlsx'), data_only=True)
M = wb['Model']; B = wb['MBR Bridge']
g = lambda a: M[a].value
REV27, E27, DA27, INT27, SBC27, SH27, MBR27 = g('Q8'), g('Q14'), g('Q17'), g('Q18'), g('Q19'), g('Q24'), g('Q10')
REV28, E28, SBC28 = g('R8'), g('R16'), g('R19')
HIT28 = -g('R15')
bridge = {c: B.cell(r, col).value for r in range(4, 14) for col, c in ((2, 'bear'), (3, 'base'), (4, 'bull')) if B.cell(r, 1).value == '2027E MBR'}

rows = list(csv.DictReader(open(D('peers_consensus_2026_10.csv'))))
peers = [r for r in rows if r['ticker'] != 'ALHC']
reg = {r['item']: float(r['value']) for r in csv.DictReader(open(D('regulatory_capital.csv')))}
tk = {r['item']: float(r['value']) for r in csv.DictReader(open(D('takeout_precedent.csv')))}

PRICE, SH_NOW, TAX, NOL, STARS = 8.28, 215.024483, 0.24, 0.463, 0.937
RATE = 0.10
STARS_2028 = HIT28 / (1 + RATE) ** (2028.5 - 2026.79) / SH_NOW
CASH, DEBT, PARENT, DIVD = reg['consolidated_cash_and_st_inv_2026_06_30_musd'], reg['convertible_face_musd'], reg['parent_cash_2026_06_30_musd'], reg['dividendable_2025_12_31_musd']
ND = {'consolidated (same as peers)': DEBT - CASH, 'parent plus dividend capacity': DEBT - PARENT - DIVD, 'parent only': DEBT - PARENT}

ev_ebitda = {r['ticker']: (float(r['price_2026_10_16']) * float(r['shares_m']) + float(r['net_debt_2025_musd'])) / float(r['ebitda_2027_cons_musd']) for r in peers}
pe = {r['ticker']: float(r['price_2026_10_16']) / float(r['adj_eps_2027_yahoo']) for r in peers}
MEV, MPE = float(np.median(list(ev_ebitda.values()))), float(np.median(list(pe.values())))
P('Peers, 2027: EV/EBITDA on S&P consensus EBITDA and consolidated net debt; P/E on Yahoo adjusted EPS')
for r in peers:
    t = r['ticker']; gr = float(r['eps_2028_ms']) / float(r['eps_2027_ms']) - 1
    P(f"  {t:4s} EV/EBITDA {ev_ebitda[t]:5.1f}x   P/E {pe[t]:5.1f}x   2028 EPS growth {gr:+.0%}")
P(f"  Median EV/EBITDA {MEV:.2f}x, median P/E {MPE:.2f}x")

alhc = [r for r in rows if r['ticker'] == 'ALHC'][0]
P(f"\nAlignment 2027 in my model: revenue ${REV27:,.0f}m, MBR {MBR27:.2%}, adj. EBITDA ${E27:,.1f}m (consensus ${float(alhc['ebitda_2027_cons_musd']):,.1f}m), stock pay ${SBC27:,.1f}m")

def legs(E, sbc_expensed=True, nd=DEBT - CASH, sh27=SH27):
    e = E - SBC27 if sbc_expensed else E
    eps = (e - DA27 - INT27) * (1 - TAX) / sh27
    return MPE * eps + NOL, (MEV * e - nd) / SH_NOW, eps

def dcf(e27, mstar=0.03, w=RATE, gT=0.03, excess=PARENT + DIVD, growth=(0.14, 0.12, 0.10, 0.08, 0.06, 0.05, 0.04, 0.03), shift=None):
    """FCF = EBITDA after stock pay, less 24% tax on EBITDA less D&A, less capex (= D&A = 0.5% of revenue).
    Regulatory capital growth and the claims float are assumed to offset. 2028 includes the stars hit."""
    m27 = (e27 - SBC27) / REV27
    m28 = (E28 - SBC28) / REV28 + (m27 - (E27 - SBC27) / REV27)
    revs = {2027: REV27, 2028: REV28}; y = 2028
    for gg in growth:
        y += 1; revs[y] = revs[y - 1] * (1 + gg)
    pv = 0.0
    for Y in range(2027, 2037):
        m = m27 if Y == 2027 else m28 if Y == 2028 else (m28 + (mstar - m28) * (Y - 2028) / 5 if Y <= 2033 else mstar)
        R = revs[Y]; eb = m * R
        fcf = eb - TAX * (eb - 0.005 * R) - 0.005 * R
        pv += fcf / (1 + w) ** (Y - 2026.29)
    tv = fcf * (1 + gT) / (w - gT) / (1 + w) ** (2036 - 2026.29)
    return (pv + tv + excess - DEBT) / SH_NOW + NOL, tv / (pv + tv)

def target(e27, sbc=True, nd=DEBT - CASH, mstar=0.03, growth=None):
    pel, evl, _ = legs(e27, sbc, nd)
    d = dcf(e27, mstar) if growth is None else dcf(e27, mstar, growth=growth)
    parts = (pel - STARS, evl - STARS, d[0] - (STARS - STARS_2028))
    return float(np.mean(parts)), parts

P('\nMatrix, before the stars cost (P/E leg | EV/EBITDA leg by cash treatment)')
for lab, s in (('Company-adjusted (stock pay left out)', False), ('Like peers (stock pay expensed)', True)):
    pel, _, eps = legs(E27, s)
    P(f"  {lab}: taxed EPS ${eps:.3f}; P/E leg ${pel:.2f}; EV/EBITDA leg " + '; '.join(f"{k} ${legs(E27, s, n)[1]:.2f}" for k, n in ND.items()))

P('\nDCF (stock pay expensed, 10%, terminal growth 3%, excess cash = parent + dividend capacity), before stars')
for w in (0.09, 0.10, 0.11):
    P(f"  WACC {w:.0%}: " + '  '.join(f"{m:.1%} ${dcf(E27, m, w)[0]:.2f}" for m in (0.02, 0.025, 0.03, 0.035, 0.04, 0.05)))
slow = (0.08, 0.06, 0.05, 0.04, 0.03, 0.03, 0.03, 0.03)
for w in (0.09, 0.10, 0.11):
    base_m = brentq(lambda m: dcf(E27, m, w)[0] - (STARS - STARS_2028) - PRICE, 0, 0.2)
    slow_m = brentq(lambda m: dcf(E27, m, w, growth=slow)[0] - (STARS - STARS_2028) - PRICE, 0, 0.2)
    P(f"  Price-implied long-run EBITDA margin after stock pay at {w:.0%}: {base_m:.2%} on my growth path, {slow_m:.2%} on the slow path")
P(f"  Margin after stock pay: 2026E {(g('P14') - g('P19')) / g('P8'):.2%}, 2027E {(E27 - SBC27) / REV27:.2%}, 2028E {(E28 - SBC28) / REV28:.2%}; terminal value share at 3% {dcf(E27)[1]:.0%}")

T, parts = target(E27)
P(f"\nBase target: P/E ${parts[0]:.2f}, EV/EBITDA ${parts[1]:.2f}, DCF ${parts[2]:.2f} -> ${T:.2f} ({T / PRICE - 1:+.0%})")
P(f"  Stars cost ${STARS:.3f}, of which the 2028 hit inside the DCF ${STARS_2028:.3f}")
alt = {}
P('  Target grid after stars (rows: earnings basis; columns: Alignment cash treatment; DCF unchanged)')
for lab, sb in (('stock pay expensed', True), ('stock pay left out', False)):
    cells = []
    for k, n in ND.items():
        alt[f'{lab}, {k}'] = target(E27, sb, n)[0]; cells.append(f"{k} ${alt[f'{lab}, {k}']:.2f}")
    P(f"    {lab}: " + '; '.join(cells))
for lab, kw in (('DCF terminal margin 4%', dict(mstar=0.04)), ('DCF terminal margin 2.5%', dict(mstar=0.025)), ('DCF slow growth', dict(growth=slow))):
    alt[lab] = target(E27, **kw)[0]
    P(f"  Sensitivity, {lab}: ${alt[lab]:.2f}")

gx = np.array([float(r['eps_2028_ms']) / float(r['eps_2027_ms']) - 1 for r in peers]); gy = np.array([pe[r['ticker']] for r in peers])
b1, b0 = np.polyfit(gx, gy, 1)
g_alhc = (g('R20')) / (g('Q20')) - 1
g_cons = float(alhc['eps_2028_ms']) / float(alhc['eps_2027_ms']) - 1
P(f"\nGrowth-adjusted P/E: across peers P/E = {b0:.1f} + {b1:.1f} x 2028 EPS growth (r = {np.corrcoef(gx, gy)[0,1]:.2f})")
P(f"  Alignment 2028 growth: my pre-tax {g_alhc:+.0%} (after the stars hit), consensus GAAP EPS {g_cons:+.0%} -> fitted P/E {b0 + b1 * g_alhc:.1f}x / {b0 + b1 * g_cons:.1f}x vs the {MPE:.1f}x median")
P(f"  Revenue growth 2027: Alignment {REV27 / g('P8') - 1:+.0%}")

shift = lambda mbr: E27 - (mbr - MBR27) * REV27
sc = {k: MBR27 + (v - bridge['base']) for k, v in bridge.items()}
vals = {k: target(shift(v))[0] for k, v in sc.items()}
pw = 0.25 * vals['bear'] + 0.5 * vals['base'] + 0.25 * vals['bull']
P(f"\nScenarios (bridge deltas on the model's {MBR27:.2%}): " + ', '.join(f"{k} {sc[k]:.2%} ${vals[k]:.2f}" for k in ('bear', 'base', 'bull')) + f"; 25/50/25 weighted ${pw:.2f}")
per_pt = (target(shift(MBR27 - 0.005))[0] - target(shift(MBR27 + 0.005))[0])
P(f"  One point of 2027 MBR is worth ${per_pt:.2f} a share")
imp = brentq(lambda m: target(shift(m))[0] - PRICE, 0.80, 0.95)
P(f"  Price-implied 2027 MBR on this method: {imp:.2%} (my {MBR27:.2%}, gap {100 * (MBR27 - imp):.2f} pts)")

# 14 Sep: same method on that day's peer prices (consensus as of today; September consensus not available),
# with 31% of the stars cost priced in (value_v8.py prior)
sep = {r['ticker']: float(r['close_2026_09_14']) for r in csv.DictReader(open(D('peers_2026_10_16.csv')))}
MEV_S = float(np.median([(sep[r['ticker']] * float(r['shares_m']) + float(r['net_debt_2025_musd'])) / float(r['ebitda_2027_cons_musd']) for r in peers]))
MPE_S = float(np.median([sep[r['ticker']] / float(r['adj_eps_2027_yahoo']) for r in peers]))
def target_sep(e27):
    e = e27 - SBC27; eps = (e - DA27 - INT27) * (1 - TAX) / SH27
    parts = (MPE_S * eps + NOL - 0.31 * STARS, (MEV_S * e - (DEBT - CASH)) / SH_NOW - 0.31 * STARS, dcf(e27)[0] - 0.31 * (STARS - STARS_2028))
    return float(np.mean(parts))
imp_sep = brentq(lambda m: target_sep(shift(m)) - 12.94, 0.70, 0.95)
P(f"  14 Sep ($12.94; peer medians {MPE_S:.1f}x P/E, {MEV_S:.1f}x EV/EBITDA): price-implied 2027 MBR {imp_sep:.2%}; change to now {100 * (imp - imp_sep):+.2f} pts")

pm = tk['bright_ca_ma_gross_price_musd'] * 1e6 / tk['bright_ca_ma_members_at_close']
evr = tk['bright_ca_ma_gross_price_musd'] / tk['bright_ca_ma_premium_revenue_2022_musd']
mem26, rev26 = g('P5') * 1000, g('P8')
tm = (pm * mem26 / 1e6 + PARENT - DEBT) / SH_NOW
trv = (evr * rev26 + PARENT - DEBT) / SH_NOW
cur_pm = (PRICE * SH_NOW + DEBT - PARENT) * 1e6 / mem26
P(f"\nTakeout check: Molina paid ${pm:,.0f} a member and {evr:.2f}x revenue for Bright's California MA plans (Jan 2024)")
P(f"  Applied to Alignment (operating plans, plus parent cash, less the notes): ${tm:.2f} a share per member, ${trv:.2f} on revenue")
P(f"  Alignment today, ex parent cash: ${cur_pm:,.0f} a member")

P(f"\nCapital (10-K FY2025): regulatory capital ${reg['regulatory_capital_2025_12_31_musd']:.1f}m vs minimum ${reg['min_requirement_2025_12_31_musd']:.1f}m "
  f"({reg['regulatory_capital_2025_12_31_musd'] / reg['min_requirement_2025_12_31_musd']:.2f}x; 2024 {reg['regulatory_capital_2024_12_31_musd'] / reg['min_requirement_2024_12_31_musd']:.2f}x); "
  f"excess ${reg['regulatory_capital_2025_12_31_musd'] - reg['min_requirement_2025_12_31_musd']:.1f}m; dividendable ${DIVD:.1f}m")
req_ratio = reg['min_requirement_2025_12_31_musd'] / reg['revenue_2025_musd']
P(f"  Minimum as a share of revenue {req_ratio:.2%}; scaled: 2026E ${req_ratio * rev26:.0f}m, 2027E ${req_ratio * REV27:.0f}m, 2028E ${req_ratio * REV28:.0f}m")
P(f"  Consolidated cash ${CASH:.1f}m vs claims payable ${reg['medical_expenses_payable_2026_06_30_musd']:.1f}m; parent ${PARENT:.1f}m vs notes ${DEBT:.0f}m")

with open(os.path.join(HERE, '..', 'results', 'valuation_bank.csv'), 'w', newline='') as fh:
    w = csv.writer(fh); w.writerow(['line', 'value'])
    for k, v in (('pe_leg', parts[0]), ('ev_leg', parts[1]), ('dcf_leg', parts[2]), ('target', T), ('prob_weighted', pw), ('implied_mbr', imp), ('per_point', per_pt)) + tuple(('alt: ' + k, v) for k, v in alt.items()):
        w.writerow([k, round(v, 4)])
open(os.path.join(HERE, '..', 'results', 'valuation_bank.txt'), 'w').write('\n'.join(out) + '\n')
