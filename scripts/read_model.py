"""Prints the key outputs of model/ALHC_model.xlsx (after recalculation) to results/model_outputs.txt.
Recalculate first: python <xlsx skill>/scripts/recalc.py model/ALHC_model.xlsx  (LibreOffice)."""
import os
from openpyxl import load_workbook
HERE = os.path.dirname(os.path.abspath(__file__))
wb = load_workbook(os.path.join(HERE, '..', 'model', 'ALHC_model.xlsx'), data_only=True)
M, V, B, S = wb['Model'], wb['Valuation'], wb['MBR Bridge'], wb['Statutory Capital']
out = []
P = lambda s: (print(s), out.append(s))
cols = {"Q3'26E": 5, "Q4'26E": 6, 'FY26E': 16, 'FY27E': 17, 'FY28E': 18}
lines = [(8, 'Revenue $m', '{:,.0f}'), (10, 'Adj. MBR', '{:.1%}'), (16, 'Adj. EBITDA $m', '{:,.1f}'), (25, 'Adj. EPS', '{:.2f}'), (22, 'Net income $m', '{:,.1f}'), (5, 'Ending members k', '{:,.1f}')]
P('Estimates      ' + '  '.join(f'{k:>8s}' for k in cols))
for r, lab, f in lines:
    P(f'{lab:15s}' + '  '.join(f'{f.format(M.cell(r, c).value):>8s}' if M.cell(r, c).value is not None else f'{"":>8s}' for c in cols.values()))
P(f"Cash end FY26/27/28: {M['P29'].value:,.0f} / {M['Q29'].value:,.0f} / {M['R29'].value:,.0f}; CFO FY27 {M['Q55'].value:,.0f}; balance check {max(abs(M.cell(44, c).value or 0) for c in range(4, 15))}")
P('\n2027 MBR bridge (bear / base / bull)')
for r in range(4, 13):
    v = [B.cell(r, c).value for c in (2, 3, 4)]
    P(f"  {B.cell(r, 1).value:32s} " + ' / '.join(f'{x:.2%}' for x in v))
P('\nCapital check ($m): FY26 / FY27 / FY28')
for r in range(5, 17):
    fmt = '{:,.2f}x' if r == 16 else '{:,.0f}'
    P(f"  {S.cell(r, 1).value:48s} " + ' / '.join(fmt.format(S.cell(r, c).value) for c in (2, 3, 4)))
P('\nValuation')
for r in range(16, 32):
    P(f"  {V.cell(r, 1).value:34s} {V.cell(r, 2).value:,.3f}")
P(f"  Peer medians: fwd P/E {V['I11'].value:.1f}x (ex-trough {V['I12'].value:.1f}x); fwd EV/EBITDA {V['K11'].value:.1f}x; at 14 Sep {V['L11'].value:.1f}x / {V['M11'].value:.1f}x")
P(f"  Price-implied 2027 MBR: now {V['B34'].value:.2%}, 14 Sep {V['B35'].value:.2%}, change {V['B36'].value*100:+.2f} pts")
for r in (50, 51, 52):
    P(f"  {V.cell(r,1).value}: MBR {V.cell(r,2).value:.2%}, EBITDA {V.cell(r,3).value:,.0f}, EPS {V.cell(r,4).value:.2f}, target ${V.cell(r,5).value:.2f}")
P('  Grid (P/E rows x MBR cols): ' + ' | '.join(f"{V.cell(41+i,1).value:.1f}x: " + ' '.join(f'{V.cell(41+i,2+j).value:.2f}' for j in range(6)) for i in range(5)))
open(os.path.join(HERE, '..', 'results', 'model_outputs.txt'), 'w').write('\n'.join(out) + '\n')
